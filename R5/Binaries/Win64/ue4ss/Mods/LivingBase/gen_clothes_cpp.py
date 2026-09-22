# gen_clothes_cpp.py -- (2026-09-14) emits the C++ item-name arrays CustomMenu.cpp's Clothes
# dropdowns need, from the SAME Other/Hair_And_Clothes_Export.xlsx source gen_clothes_lua.py reads.
# The DLL can't read config.lua at runtime (same situation as kHairFriendlyNames/kClothColors,
# see CustomMenu.cpp's own comments on those) -- these arrays are generated once and pasted into
# CustomMenu.cpp by hand, then KEPT IN SYNC BY HAND if the spreadsheet is ever revised (re-run this
# and re-paste).
#
# Run this after ANY edit to Hair_And_Clothes_Export.xlsx that should reach the Clothes menu, then
# paste the printed block over the existing kClothes*Items/kClothesOutfits arrays in CustomMenu.cpp.
#
# Usage:  python gen_clothes_cpp.py

import openpyxl, os, re, shutil, sys
from collections import defaultdict

XLSX_PATH = os.environ.get("CLOTHES_XLSX_OVERRIDE") or r"H:\OneDrive\Coding\WINDROSE MODS\Other\Hair_And_Clothes_Export.xlsx"

# Mask ADDED (2026-09-16, RedFalcon: "add it back in... treat it the same as waist or cape") --
# mirrors gen_clothes_lua.py's own REAL_BODY_PARTS change. Unlike that script, no OUTFIT_BODY_PARTS
# split is needed here: this generator only counts WHICH body parts a Clothing-Outfits set has (to
# decide if it qualifies for the dropdown, REQUIRED_OUTFIT_PARTS) -- it never resolves or emits any
# actual per-slot outfit PIECE data (that only happens server-side, in Lua, off
# Config.CLOTHES_OUTFITS, which gen_clothes_lua.py's own OUTFIT_BODY_PARTS already keeps Mask out
# of) -- so Mask counting toward a set's membership here has no way to make Mask part of any
# outfit's actual applied pieces.
REAL_BODY_PARTS = {"Torso", "Legs", "Waist", "Hands", "Feets", "Headgear", "Cape", "Mask"}
REQUIRED_OUTFIT_PARTS = {"Feets", "Torso", "Legs"}

# Maps a real body-part key to this generator's own C++ array-name suffix (matches the Custom tab's
# own friendlier row labels for Feet/Hat -- Feets->Feet, Headgear->Hat -- everything else unchanged).
CPP_SUFFIX = {
    "Torso": "Torso", "Legs": "Legs", "Waist": "Waist", "Hands": "Hands",
    "Feets": "Feet", "Headgear": "Hat", "Cape": "Cape", "Mask": "Mask",
}


def open_workbook(path):
    try:
        return openpyxl.load_workbook(path, data_only=True)
    except PermissionError:
        fallback = os.path.join(os.path.dirname(path), "Hair_And_Clothes_Export_copy.xlsx")
        try:
            shutil.copy(path, fallback)
        except Exception:
            print(f"ERROR: '{path}' is locked (probably open in Excel) and no fallback copy "
                  f"could be made. Close Excel and re-run.", file=sys.stderr)
            raise
        print(f"NOTE: source was locked, read a fresh copy at '{fallback}' instead.")
        return openpyxl.load_workbook(fallback, data_only=True)


def norm_ws(s):
    if s is None:
        return None
    return re.sub(r"\s+", " ", str(s)).strip()


def cppstr(s):
    s = str(s).replace("\\", "\\\\").replace('"', '\\"')
    return '"' + s + '"'


wb = open_workbook(XLSX_PATH)

# ---- Clothes Adjusted -> per-slot item lists (name + locked) ----
ws = wb["Clothes Adjusted"]
items_by_part = defaultdict(list)
seen_friendly = defaultdict(int)

for row in ws.iter_rows(min_row=2, values_only=True):
    bodyPart, name, friendlyName, setName, availability, maleMesh, femaleMesh, unisexMesh, _src = row[:9]
    bodyPart = norm_ws(bodyPart)
    if not bodyPart or bodyPart not in REAL_BODY_PARTS:
        continue
    if not (maleMesh or femaleMesh or unisexMesh):
        continue
    friendlyName = norm_ws(friendlyName)
    availability = norm_ws(availability) or ""

    key = (bodyPart, friendlyName)
    seen_friendly[key] += 1
    finalFriendly = friendlyName
    if seen_friendly[key] > 1:
        finalFriendly = f"{friendlyName} ({seen_friendly[key]})"

    locked = "only with unlock" in availability.lower()
    items_by_part[bodyPart].append({"friendlyName": finalFriendly, "locked": locked})

# ---- Mask supplement: pulled from the raw "Clothes" sheet, not "Clothes Adjusted" (2026-09-16) --
# same reasoning as gen_clothes_lua.py's own identical supplement (Mask has no rows in "Clothes
# Adjusted", RedFalcon's own curated pass never covered it since Mask was excluded when that sheet
# was built).
wsRaw = wb["Clothes"]
for row in wsRaw.iter_rows(min_row=2, values_only=True):
    bodyPart, name, friendlyName, availability, maleMesh, femaleMesh, unisexMesh, _src = row[:8]
    bodyPart = norm_ws(bodyPart)
    if bodyPart != "Mask":
        continue
    if not (maleMesh or femaleMesh or unisexMesh):
        continue
    friendlyName = norm_ws(friendlyName)
    availability = norm_ws(availability) or ""

    key = (bodyPart, friendlyName)
    seen_friendly[key] += 1
    finalFriendly = friendlyName
    if seen_friendly[key] > 1:
        finalFriendly = f"{friendlyName} ({seen_friendly[key]})"

    locked = "only with unlock" in availability.lower()
    items_by_part[bodyPart].append({"friendlyName": finalFriendly, "locked": locked})

out = []
out.append("// Clothes item lists (2026-09-14) -- generated by gen_clothes_cpp.py from the SAME")
out.append("// Other/Hair_And_Clothes_Export.xlsx source Config.CLOTHES_ITEMS (config.lua) is built")
out.append("// from -- see gen_clothes_lua.py. KEEP IN SYNC BY HAND with config.lua if the spreadsheet")
out.append("// is ever revised (re-run gen_clothes_cpp.py, re-paste). `locked` items (\"Only with")
out.append("// Unlock\" in the sheet) are skipped entirely from the dropdown unless g_clothesUnlocked")
out.append("// is true (mirrors Config.CLOTHES_UNLOCK_ALL, read live from clothes_unlock_state.txt --")
out.append("// see pollClothesUnlockState). Index 0 in every array is the reserved \"(Remove)\" sentinel")
out.append("// (RedFalcon: \"add a remove to each body part at the top of the list\") -- never locked.")
out.append("struct ClothesItem { const char* name; bool locked; };")
for bodyPart in ["Torso", "Legs", "Waist", "Hands", "Feets", "Headgear", "Cape", "Mask"]:
    rows = sorted(items_by_part[bodyPart], key=lambda r: r["friendlyName"].lower())
    suffix = CPP_SUFFIX[bodyPart]
    out.append(f"constexpr ClothesItem kClothes{suffix}Items[] = {{")
    out.append('    { "(Remove)", false },')
    for r in rows:
        out.append(f'    {{ {cppstr(r["friendlyName"])}, {"true" if r["locked"] else "false"} }},')
    out.append("};")
out.append("")

# ---- Clothing Outfits -> outfit names (filtered to qualifying sets, same rule as gen_clothes_lua.py) ----
ws2 = wb["Clothing Outfits"]
pieces_by_set = defaultdict(set)
set_order = []
for row in ws2.iter_rows(min_row=2, values_only=True):
    bodyPart, name, friendlyName, setName = row[:4]
    bodyPart = norm_ws(bodyPart)
    setName = norm_ws(setName)
    if not setName or not bodyPart or bodyPart not in REAL_BODY_PARTS:
        continue
    if setName not in pieces_by_set:
        set_order.append(setName)
    pieces_by_set[setName].add(bodyPart)

qualifying = sorted([s for s in set_order if REQUIRED_OUTFIT_PARTS.issubset(pieces_by_set[s])], key=lambda s: s.lower())

out.append("// Outfit dropdown (top of the Clothes section) -- same qualifying-set filter as")
out.append("// gen_clothes_lua.py's Config.CLOTHES_OUTFITS (must define at least Feet+Torso+Legs).")
out.append("// Index 0 is the reserved \"(Remove All)\" sentinel (RedFalcon: \"a remove all on the")
out.append("// outfits\"). No `locked` concept here -- an outfit referencing a currently-locked piece")
out.append("// just silently skips that one piece on apply (Spawner.ApplyClothesOutfit's own per-piece")
out.append("// Spawner.ApplyClothesItem call already handles that gracefully).")
out.append("constexpr const char* kClothesOutfitNames[] = {")
out.append('    "(Remove All)",')
for s in qualifying:
    out.append(f"    {cppstr(s)},")
out.append("};")

print("\n".join(out))
